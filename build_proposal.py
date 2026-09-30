from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = 'inimu_企画提案書_FINAL_2026-09-30.docx'
NAVY = '243746'; MUTED = '66737C'; ACCENT = 'A36F4F'; PALE = 'F5F1EC'; LINE = 'D9D3CC'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(cell, color=LINE, size='6'):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = b.find(qn('w:'+edge))
        if e is None: e=OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),size); e.set(qn('w:color'),color)

def run_font(run, size=11, color=NAVY, bold=False, italic=False):
    run.font.name='Yu Gothic'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic')
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color); run.bold=bold; run.italic=italic

def para(doc, text='', size=11, color=NAVY, bold=False, italic=False, align=None, before=0, after=8):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.18
    if align is not None: p.alignment=align
    r=p.add_run(text); run_font(r,size,color,bold,italic); return p

def heading(doc, text, level=1):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(16 if level==1 else 10); p.paragraph_format.space_after=Pt(7); p.paragraph_format.keep_with_next=True
    r=p.add_run(text); run_font(r, 20 if level==1 else 14, NAVY, True); return p

def bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.15
        r=p.add_run(item); run_font(r,10.5,NAVY)

def numbered(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.15
        r=p.add_run(item); run_font(r,10.5,NAVY)

def callout(doc, label, text):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(6.25)
    c=t.cell(0,0); shade(c,PALE); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(label); run_font(r,9,ACCENT,True)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(8); r=p.add_run(text); run_font(r,11,NAVY)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def page_title(doc, no, title, subtitle=''):
    para(doc, f'{no:02d}  /  {title}', size=9, color=ACCENT, bold=True, after=10)
    para(doc, title, size=25, color=NAVY, bold=True, after=3)
    if subtitle: para(doc, subtitle, size=12, color=MUTED, after=20)

def footer(section):
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('inimu  |  企画提案書 FINAL'); run_font(r,8,MUTED)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
for s in doc.styles:
    if s.type==WD_STYLE_TYPE.PARAGRAPH:
        s.font.name='Yu Gothic'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic')
normal=doc.styles['Normal']; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(NAVY)
for name in ('List Bullet','List Number'):
    st=doc.styles[name]; st.font.name='Yu Gothic'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic'); st.font.size=Pt(10.5)
footer(sec)

# Cover
para(doc,'CLIENT WORK',size=10,color=ACCENT,bold=True,after=46)
para(doc,'inimu',size=14,color=MUTED,bold=True,after=8)
para(doc,'Web Experience Redesign',size=31,color=NAVY,bold=True,after=8)
para(doc,'学校提出用 企画提案書 FINAL版',size=15,color=MUTED,after=26)
callout(doc,'企画の核','情報を増やすのではなく、inimuらしさが伝わる順番に情報を編集する。')
para(doc,'提出日  2026.09.30',size=10,color=MUTED,after=4)
para(doc,'対象：inimu Webサイト / SHOP / EXPERIENCE / 会員導線',size=10,color=MUTED,after=8)
doc.add_page_break()

# Concept
page_title(doc,1,'CONCEPT','情報設計を、ブランド体験の設計へ')
heading(doc,'背景',1)
para(doc,'inimuには商品、ワークショップ、素材、土地、職人、ブランドストーリーなど、伝える価値が多くあります。一方で情報量が多いほど、初めて訪れた人は「何がinimuなのか」をつかむ前にページを離れてしまう可能性があります。')
heading(doc,'提案の考え方',1)
para(doc,'情報を増やすのではなく、目的と利用意図に合わせて情報を再編集します。入口を整理し、ブランドの核を先に見せ、その後に商品・体験・会員・調査へ自然につなげます。')
callout(doc,'仮説','ユーザーが最初に理解すべきなのは、商品の数ではなく、inimuが「香りを買う場所」であり「香りをつくる場所」でもあること。')
heading(doc,'体験フロー',1)
numbered(doc,['TOPで目的を選ぶ：SHOP / EXPERIENCE','SHOPで世界観を知る：SIGNATURE STORIES → 香り・アイテム・ブランド','商品やブランドを見たあとに、自分でつくる選択肢へ進む','EXPERIENCEで短時間のワークショップを体験する','会員情報・アンケートを次の改善へつなげる'])
doc.add_page_break()

# Cause
page_title(doc,2,'01 CAUSE / 課題','見えない課題を、行動の詰まりとして捉える')
heading(doc,'観察されたこと',1)
bullets(doc,['商品・体験・ブランドストーリーなどの情報量が多い','必要な情報が複数ページ・階層に分散している','SHOPとEXPERIENCEのつながりが弱く、次の行動が選びにくい','初めて来た人がブランド背景まで到達しにくい','一般的な香り表現だけでは、inimu独自の価値が伝わりにくい'])
heading(doc,'口コミ・反応を読む視点',1)
para(doc,'短い感想の数だけで評価せず、「何が書かれているか」「どこで理解が止まったか」「何が記憶に残ったか」を読みます。評価点や件数は結果の強弱を示しますが、改善の方向は具体的な言葉から考えます。')
callout(doc,'判断の前提','以下は現時点の企画仮説であり、実装後のWeb計測・アンケート・インタビューで検証します。')
doc.add_page_break()

# Hypothesis
page_title(doc,3,'02 HYPOTHESIS / 仮説','「情報不足」ではなく「意味づけ不足」')
heading(doc,'仮説A｜情報設計',1)
para(doc,'必要な情報がないのではなく、必要な情報が必要な場所につながっていない。そのため、最初にブランドの核を見せ、次に選択肢を分ける構造へ変える。')
heading(doc,'仮説B｜体験価値',1)
para(doc,'inimuの固有価値は、香り単体ではなく、土地・素材・文化・職人の背景と香りを一緒に体験できることにある。')
heading(doc,'仮説C｜会員導線',1)
para(doc,'会員登録は目的ではなく、体験後の記憶を保存し、次の提案につなげるための接点として設計する。')
heading(doc,'仮説D｜ブランド',1)
para(doc,'HATENKO / WANOWAは、商品一覧の一部ではなく、inimuを特徴づける二つの入口として前面に出す。')
doc.add_page_break()

# Idea 1
page_title(doc,4,'03 IDEA (1)','Web / 情報設計')
heading(doc,'TOP｜目的別の入口',1)
bullets(doc,['SHOP：香りを買う。','EXPERIENCE：香りをつくる。'])
para(doc,'最初に利用目的を選べるようにし、ユーザーが自分の関心に合う場所へ迷わず進めるようにします。')
heading(doc,'SHOP｜ブランドの核を先に見せる',1)
callout(doc,'SIGNATURE STORIES','inimuをかたちづくる、二つの香り。\n01 HATENKO｜伝統と革新で、型を破る香り。\n02 WANOWA｜土地と素材から生まれる香り。')
para(doc,'その後に「香りから選ぶ」「アイテムから探す」「ブランドから探す」を配置し、理解してから選べる順番にします。ページ最下部にはMAKE YOUR SCENTを置き、購入後の自然な体験導線として扱います。')
heading(doc,'EXPERIENCE｜短時間で理解できる体験',1)
para(doc,'香りを知る → 選ぶ → つくる → 持ち帰る、という流れを明確にし、FAQやアクセス情報まで一つの体験として整理します。')
doc.add_page_break()

# Idea 2
page_title(doc,5,'03 IDEA (2)','アンケート / 顧客理解')
callout(doc,'30 SEC / 4-TAP SURVEY','体験後に、短時間で答えられる選択式アンケートを設置します。')
heading(doc,'質問設計',1)
bullets(doc,['何がきっかけでinimuを選んだか','一番印象に残ったこと','他のお店との違いを感じたこと','購入または体験後の期待','ブランドへの関心'])
heading(doc,'自由記述の位置づけ',1)
para(doc,'自由記述は任意とし、選択式の回答を中心にします。回答負担を抑えながら、「何が価値だったか」を比較できるデータにします。')
heading(doc,'画面イメージ',1)
numbered(doc,['体験を選ぶ','四つの選択肢から一つ選ぶ','理由を任意で入力','送信して完了'])
doc.add_page_break()

# Verification
page_title(doc,6,'04 VERIFICATION / 検証','評価点ではなく、行動と発言を組み合わせる')
heading(doc,'Webで測定するもの',1)
bullets(doc,['HATENKO / WANOWA STORYの閲覧','STORYから商品詳細への遷移','SHOPからEXPERIENCEへの遷移','EXPERIENCEの予約・問い合わせ','ページ滞在時間と離脱箇所'])
heading(doc,'アンケートで聞くこと',1)
para(doc,'Webの行動データと、ユーザーの言葉を組み合わせて評価します。クリックが増えたかだけでなく、ブランドを理解したうえで次の行動を選べたかを見ます。')
callout(doc,'検証の問い','「SIGNATURE STORIESを先に見せると、inimuの違いが理解され、商品・体験への遷移が増えるか」')
heading(doc,'レビュー方法',1)
para(doc,'定量データは傾向を、自由記述・インタビューは理由を示します。両者が一致しない場合も、仮説を見直す材料として扱います。')
doc.add_page_break()

# Budget
page_title(doc,7,'05 BUDGET / SCHEDULE','実装と検証を分けて、無理なく進める')
heading(doc,'作業項目',1)
rows=[('調査・情報整理','既存サイト、ブランド、導線の確認'),('情報設計','TOP / SHOP / EXPERIENCEの優先順位'),('ビジュアル設計','HATENKO / WANOWAを含む見せ方'),('実装','HTML / CSS / JS、レスポンシブ'),('会員導線','登録情報、管理、CSV'),('計測・アンケート','イベント設計、4タップ調査'),('テスト・改善','表示確認、導線確認、修正')]
t=doc.add_table(rows=1, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(1.55); t.columns[1].width=Inches(4.7)
for i,h in enumerate(['項目','内容']): c=t.rows[0].cells[i]; c.text=h; shade(c,PALE); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for a,b in rows:
    cells=t.add_row().cells
    for i,v in enumerate((a,b)): cells[i].text=v; borders(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cells[i].paragraphs:
        for r in p.runs: run_font(r,9.5,NAVY, i==0)
heading(doc,'進行イメージ',1)
numbered(doc,['要件・既存構造の確認','情報設計と画面構成の整理','実装・レスポンシブ確認','会員導線とCSVの確認','公開前テスト','アンケート・データをもとに改善'])
doc.add_page_break()

# Simulation
page_title(doc,8,'06 SIMULATION','つくって、測って、戻す')
heading(doc,'データフロー',1)
callout(doc,'登録データ','新規会員登録 → localStorage → 管理画面 → CSV書き出し → LibreOffice Calc')
heading(doc,'保存する情報',1)
bullets(doc,['会員ID / 氏名 / フリガナ / メールアドレス','電話番号 / 郵便番号 / 住所','香りに求めること / 登録日時'])
heading(doc,'改善サイクル',1)
numbered(doc,['Webの閲覧・遷移を確認','アンケートで理由を聞く','HATENKO / WANOWA、商品、体験のどこで理解が進んだかを見る','次のWeb・商品・体験企画へ反映する'])
heading(doc,'学校提出用デモの範囲',1)
para(doc,'本提案の会員登録・CSVは学校提出用デモとして扱います。実在する個人情報は入力せず、本番利用時には認証、アクセス制御、個人情報保護、サーバー側バリデーション、公開範囲管理が必要です。')
doc.add_page_break()

# Closing
page_title(doc,9,'FINAL MESSAGE','inimuらしさを、見つけやすくする')
para(doc,'調査を進める中で、HATENKOとWANOWAには、伝統、土地、素材、職人、記憶など、inimu独自の価値がすでに含まれていることが分かりました。課題は価値がないことではなく、その価値が商品一覧や別ページの奥に分散していることです。')
para(doc,'そこで本提案では、情報を増やすより先に、意味の順番を整えます。最初に二つの香りを見せ、次に選ぶ方法を示し、最後に自分でつくる可能性へつなげます。')
callout(doc,'最終メッセージ','選ぶだけでなく、つくる。\ninimuの固有の価値を、Webと顧客の声によって育てていく。')
heading(doc,'FACT / HYPOTHESIS / IDEA',1)
bullets(doc,['FACT：既存Webと既存コンテンツから確認できたこと','HYPOTHESIS：今回の調査から置いた仮説','IDEA：次のWeb・商品・体験改善へ進む提案'])
para(doc,'以上を、学校提出用の企画提案書 FINAL版としてまとめます。',size=10,color=MUTED,after=0)

doc.save(OUT)
print(OUT)
